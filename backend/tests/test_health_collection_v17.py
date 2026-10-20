"""v1.7 health-collection feature tests.

Covers:

* ``_parse_optical`` / ``_parse_config_excerpt`` pure functions
  (no DB / no SSH).
* ``GET /api/v1/health/snapshots`` list endpoint with metric filter.
* ``GET /api/v1/health/snapshots/{id}/export?format=csv|xlsx|txt``
  with the three formats and a custom filename.
* ``/api/v1/health/schedules`` CRUD + ``/run``.

The schedule ``/run`` endpoint fans out to ``health_snapshot_one.delay``;
we patch the task's ``.delay`` to a no-op so the test stays
synchronous and doesn't need a running broker.
"""
from __future__ import annotations

import csv
import io
import os
import tempfile
import unittest
from datetime import datetime, timezone

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from tests.test_jobs_dry_run_e2e import (  # type: ignore  # noqa: E402
    _make_test_app, _mint_token, _create_user, _create_simulated_device,
)


# ── Parser unit tests ────────────────────────────────────────────────

from app.core.enums import Vendor
from app.services.health_collector import (
    _parse_optical, _parse_config_excerpt,
)


class OpticalParserTest(unittest.TestCase):
    def test_huawei_verbose_rows(self) -> None:
        text = (
            "GigabitEthernet0/0/1  -7.23  -5.10  -7.00  -1.00  normal\n"
            "GigabitEthernet0/0/2  -8.50  -6.20  -7.00  -1.00  normal\n"
        )
        out = _parse_optical(Vendor.HUAWEI_VRP, text)
        self.assertEqual(len(out), 2)
        self.assertEqual(out[0]["port"], "GigabitEthernet0/0/1")
        self.assertAlmostEqual(out[0]["rx_dbm"], -7.23)
        self.assertAlmostEqual(out[0]["tx_dbm"], -5.10)
        self.assertEqual(out[1]["port"], "GigabitEthernet0/0/2")

    def test_empty_text(self) -> None:
        self.assertEqual(_parse_optical(Vendor.HUAWEI_VRP, ""), [])
        self.assertEqual(_parse_optical(Vendor.CISCO_IOS_XE, "no DDM support here"), [])

    def test_garbage_returns_empty(self) -> None:
        # Text that has a port-like token but no two consecutive
        # signed dBm numbers.  The regex won't find a match, so the
        # parser returns an empty list rather than guessing.
        text = "GigabitEthernet0/0/1  some text but no values\n"
        out = _parse_optical(Vendor.HUAWEI_VRP, text)
        self.assertEqual(out, [])


class ConfigFingerprintTest(unittest.TestCase):
    def test_basic(self) -> None:
        out = _parse_config_excerpt(Vendor.HUAWEI_VRP, "line1\nline2\nline3\n")
        self.assertEqual(out["lines"], 3)
        self.assertTrue(out["fingerprint"] and len(out["fingerprint"]) == 64)
        self.assertEqual(out["excerpt"], "line1\nline2\nline3")

    def test_empty(self) -> None:
        out = _parse_config_excerpt(Vendor.HUAWEI_VRP, "")
        self.assertEqual(out["fingerprint"], "")
        self.assertEqual(out["lines"], 0)


# ── HTTP e2e ─────────────────────────────────────────────────────────

class HealthSnapshotsE2ETest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app, cls.engine, cls.db_path = _make_test_app()
        cls.admin_id = _create_user(cls.engine)
        cls.device_id = _create_simulated_device(cls.engine)
        cls.client = TestClient(cls.app)
        cls.token = _mint_token(cls.admin_id)
        cls.headers = {"Authorization": f"Bearer {cls.token}"}

        # Seed a couple of snapshot rows directly so we can test the
        # list / export endpoints without spinning up Celery.  We
        # also want to confirm the new columns round-trip.
        from app.models.health import DeviceHealthSnapshot
        Session = sessionmaker(bind=cls.engine, future=True)
        db = Session()
        try:
            db.add(DeviceHealthSnapshot(
                device_id=cls.device_id,
                collected_at=datetime(2026, 1, 1, 10, 0, tzinfo=timezone.utc),
                vendor="huawei_vrp",
                cpu_pct=12.0, mem_pct=34.0, temp_c=38.0,
                fan_state="normal", power_w=45.0,
                optical=[
                    {"port": "GigabitEthernet0/0/1", "rx_dbm": -7.23, "tx_dbm": -5.10},
                ],
                config_fingerprint="abc123" + "0" * 58,
                config_lines=42,
                config_excerpt="!\nsysname x\nreturn\n",
                metrics_run=["cpu", "mem", "optical"],
                source="collect",
            ))
            db.add(DeviceHealthSnapshot(
                device_id=cls.device_id,
                collected_at=datetime(2026, 1, 1, 11, 0, tzinfo=timezone.utc),
                vendor="huawei_vrp",
                cpu_pct=15.0,
                metrics_run=["cpu"],
                source="collect",
            ))
            db.commit()
        finally:
            db.close()

    @classmethod
    def tearDownClass(cls) -> None:
        try:
            os.unlink(cls.db_path)
        except OSError:
            pass

    def test_list_snapshots(self) -> None:
        r = self.client.get("/api/v1/health/snapshots", headers=self.headers)
        self.assertEqual(r.status_code, 200, r.text)
        data = r.json()
        self.assertGreaterEqual(len(data), 2)
        # Newest first; pydantic may drop the trailing "+00:00" on
        # serialisation depending on the schema version, so we
        # compare on the first 19 chars (date + HH:MM:SS).
        self.assertTrue(data[0]["collected_at"].startswith("2026-01-01T11:00:00"))
        # New fields are exposed
        self.assertIn("optical", data[1])
        self.assertEqual(data[1]["config_fingerprint"], "abc123" + "0" * 58)
        self.assertEqual(data[1]["metrics_run"], ["cpu", "mem", "optical"])

    def test_list_filter_by_metric(self) -> None:
        r = self.client.get(
            "/api/v1/health/snapshots",
            params={"metric": "optical"},
            headers=self.headers,
        )
        self.assertEqual(r.status_code, 200, r.text)
        data = r.json()
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]["metrics_run"], ["cpu", "mem", "optical"])

    def test_list_filter_by_device(self) -> None:
        r = self.client.get(
            f"/api/v1/health/devices/{self.device_id}/snapshots",
            headers=self.headers,
        )
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(len(r.json()), 2)

    def test_export_csv(self) -> None:
        # Find the snapshot that has an ``optical`` payload — that's
        # the one seeded in setUpClass with a single port entry.
        r = self.client.get("/api/v1/health/snapshots", headers=self.headers)
        rows = r.json()
        target = next(
            x for x in rows
            if (x.get("optical") or [])
        )
        target_id = target["id"]
        r2 = self.client.get(
            f"/api/v1/health/snapshots/{target_id}/export",
            params={"format": "csv", "filename": "my-file"},
            headers=self.headers,
        )
        self.assertEqual(r2.status_code, 200, r2.text)
        self.assertIn("text/csv", r2.headers["content-type"])
        cdisp = r2.headers.get("content-disposition", "")
        self.assertIn("my-file-", cdisp)
        # Parse the CSV body
        body = r2.content.decode("utf-8")
        reader = csv.DictReader(io.StringIO(body))
        out_rows = list(reader)
        self.assertEqual(len(out_rows), 1)
        self.assertEqual(out_rows[0]["snapshot_id"], str(target_id))
        self.assertEqual(out_rows[0]["optical_port_count"], "1")

    def test_export_xlsx(self) -> None:
        from openpyxl import load_workbook
        r = self.client.get("/api/v1/health/snapshots", headers=self.headers)
        first_id = r.json()[0]["id"]
        r2 = self.client.get(
            f"/api/v1/health/snapshots/{first_id}/export",
            params={"format": "xlsx"},
            headers=self.headers,
        )
        self.assertEqual(r2.status_code, 200, r2.text)
        self.assertIn(
            "spreadsheetml",
            r2.headers["content-type"],
        )
        wb = load_workbook(io.BytesIO(r2.content))
        ws = wb.active
        # First row is the header; second row is the data row.
        self.assertGreaterEqual(ws.max_row, 2)
        # First column is snapshot_id
        self.assertEqual(int(ws.cell(row=2, column=1).value), first_id)

    def test_export_txt(self) -> None:
        r = self.client.get("/api/v1/health/snapshots", headers=self.headers)
        first_id = r.json()[0]["id"]
        r2 = self.client.get(
            f"/api/v1/health/snapshots/{first_id}/export",
            params={"format": "txt", "filename": "report"},
            headers=self.headers,
        )
        self.assertEqual(r2.status_code, 200, r2.text)
        self.assertIn("text/plain", r2.headers["content-type"])
        body = r2.content.decode("utf-8")
        self.assertIn("snapshot_id:", body)
        self.assertIn(str(first_id), body)

    def test_export_invalid_format_422(self) -> None:
        r = self.client.get(
            "/api/v1/health/snapshots/1/export",
            params={"format": "pdf"},
            headers=self.headers,
        )
        # Pydantic regex on the Query param -> 422 (FastAPI's
        # validation pattern rejects unknown values).
        self.assertEqual(r.status_code, 422)

    def test_export_missing_snapshot_404(self) -> None:
        r = self.client.get(
            "/api/v1/health/snapshots/999999/export",
            params={"format": "csv"},
            headers=self.headers,
        )
        self.assertEqual(r.status_code, 404)

    def test_filename_sanitised(self) -> None:
        r = self.client.get("/api/v1/health/snapshots", headers=self.headers)
        first_id = r.json()[0]["id"]
        # The sanitiser's contract is "no path separators in the
        # output".  Path-traversal (``..``) is only dangerous when
        # the OS is asked to *resolve* it as a path component —
        # once we've stripped the slashes, the remaining ``..`` is
        # just a literal 2-character substring of the filename.
        r2 = self.client.get(
            f"/api/v1/health/snapshots/{first_id}/export",
            params={"format": "csv", "filename": "../../etc/passwd"},
            headers=self.headers,
        )
        self.assertEqual(r2.status_code, 200, r2.text)
        cdisp = r2.headers.get("content-disposition", "")
        self.assertNotIn("/", cdisp)
        # And no NUL byte (would be a security problem in C
        # bindings downstream).
        self.assertNotIn("\x00", cdisp)


# ── Schedule CRUD ───────────────────────────────────────────────────

class HealthSchedulesCRUDTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app, cls.engine, cls.db_path = _make_test_app()
        cls.admin_id = _create_user(cls.engine)
        cls.device_id = _create_simulated_device(cls.engine)
        cls.client = TestClient(cls.app)
        cls.token = _mint_token(cls.admin_id)
        cls.headers = {"Authorization": f"Bearer {cls.token}"}

        # Patch health_snapshot_one.delay so /run is synchronous in
        # the test process (no broker required).
        from app.api.v1 import health_schedules as sched_mod
        from app.workers import tasks as tasks_mod

        class _FakeTask:
            def __init__(self) -> None:
                self.id = "fake-task-id"
        cls._fired: list[tuple] = []

        def _fake_delay(device_id, metrics=None):  # type: ignore[no-untyped-def]
            cls._fired.append((device_id, metrics))
            return _FakeTask()

        cls._orig_delay = tasks_mod.health_snapshot_one.delay
        tasks_mod.health_snapshot_one.delay = staticmethod(_fake_delay)  # type: ignore[assignment]

    @classmethod
    def tearDownClass(cls) -> None:
        from app.workers import tasks as tasks_mod
        tasks_mod.health_snapshot_one.delay = cls._orig_delay  # type: ignore[assignment]
        try:
            os.unlink(cls.db_path)
        except OSError:
            pass

    def setUp(self) -> None:
        type(self)._fired = []  # reset between tests

    def test_create_and_list(self) -> None:
        r = self.client.post(
            "/api/v1/health/schedules",
            json={
                "name": "core switches every 5m",
                "interval_minutes": 5,
                "device_ids": [self.device_id],
                "metrics": ["cpu", "mem", "optical"],
                "enabled": True,
            },
            headers=self.headers,
        )
        self.assertEqual(r.status_code, 201, r.text)
        body = r.json()
        self.assertEqual(body["interval_minutes"], 5)
        self.assertEqual(body["device_ids"], [self.device_id])
        self.assertEqual(body["metrics"], ["cpu", "mem", "optical"])
        self.assertTrue(body["enabled"])

        # List
        r2 = self.client.get("/api/v1/health/schedules", headers=self.headers)
        self.assertEqual(r2.status_code, 200)
        self.assertGreaterEqual(len(r2.json()), 1)

    def test_unknown_metric_rejected(self) -> None:
        r = self.client.post(
            "/api/v1/health/schedules",
            json={
                "name": "bogus",
                "interval_minutes": 5,
                "metrics": ["cpu", "made_up_metric"],
            },
            headers=self.headers,
        )
        self.assertEqual(r.status_code, 422)

    def test_interval_too_small_rejected(self) -> None:
        r = self.client.post(
            "/api/v1/health/schedules",
            json={"name": "x", "interval_minutes": 0},
            headers=self.headers,
        )
        self.assertEqual(r.status_code, 422)

    def test_run_now_dispatches_tasks(self) -> None:
        r = self.client.post(
            "/api/v1/health/schedules",
            json={
                "name": "fanout",
                "interval_minutes": 5,
                "device_ids": [self.device_id],
                "metrics": ["cpu", "optical"],
            },
            headers=self.headers,
        )
        self.assertEqual(r.status_code, 201, r.text)
        sched_id = r.json()["id"]

        r2 = self.client.post(
            f"/api/v1/health/schedules/{sched_id}/run",
            headers=self.headers,
        )
        self.assertEqual(r2.status_code, 200, r2.text)
        body = r2.json()
        # ``last_fired_at`` is "now"; ``next_run_at`` is
        # now + interval (5 min).  They should NOT be equal —
        # the manual run pushed the schedule forward by the
        # interval, so a sweep tick right after the manual
        # dispatch won't re-fire the same plan.
        self.assertNotEqual(body["last_fired_at"], body["next_run_at"])
        self.assertIsNotNone(body["last_fired_at"])
        # ``metrics`` is forwarded to the worker
        self.assertEqual(type(self)._fired, [(self.device_id, ["cpu", "optical"])])

    def test_delete(self) -> None:
        r = self.client.post(
            "/api/v1/health/schedules",
            json={"name": "to-delete", "interval_minutes": 10},
            headers=self.headers,
        )
        sched_id = r.json()["id"]
        r2 = self.client.delete(
            f"/api/v1/health/schedules/{sched_id}",
            headers=self.headers,
        )
        self.assertEqual(r2.status_code, 204)
        r3 = self.client.get(
            f"/api/v1/health/schedules/{sched_id}",
            headers=self.headers,
        )
        self.assertEqual(r3.status_code, 404)
